from __future__ import annotations

import gzip
import json
import logging
import os
import re
import tempfile
from pathlib import Path
from typing import BinaryIO
from uuid import uuid4

from ..config import Settings
from ..models import ModelGateway
from .mineru import extract_with_mineru

logger = logging.getLogger("law_rag.workspace.file_utits")


def save_json_file(path: Path, payload: dict, compressed: bool = False) -> Path:
    """将字典以 UTF-8 JSON 写入目标路径，并可选择使用 gzip 压缩。"""
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    content = json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")
    if compressed:
        target = target.with_name(target.name + ".gz")
        target.write_bytes(gzip.compress(content, compresslevel=6))
    else:
        target.write_bytes(content)
    return target


def load_json_file(path: Path) -> dict:
    """读取普通或 gzip JSON 文件；顶层不是对象时返回空字典。"""
    source = Path(path)
    content = gzip.decompress(source.read_bytes()).decode("utf-8") if source.suffix == ".gz" else source.read_text(encoding="utf-8")
    value = json.loads(content)
    return value if isinstance(value, dict) else {}
TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".rst", ".csv", ".tsv", ".log", ".json", ".xml", ".html", ".htm", ".yaml", ".yml", ".ini", ".conf", ".sql", ".js", ".ts", ".py", ".java", ".go", ".rs", ".c", ".cpp", ".h", ".hpp"}
IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff", ".webp"}
AUDIO_SUFFIXES = {".mp3", ".wav", ".m4a", ".aac", ".flac", ".ogg", ".opus"}
VIDEO_SUFFIXES = {".mp4", ".mov", ".avi", ".mkv", ".webm", ".mpeg", ".mpg"}
MEDIA_SUFFIXES = IMAGE_SUFFIXES | AUDIO_SUFFIXES | VIDEO_SUFFIXES
ALLOWED_UPLOAD_SUFFIXES = {".txt", ".md", ".markdown", ".rst", ".pdf", ".docx", ".xlsx", ".csv", ".tsv", ".log", *IMAGE_SUFFIXES, *AUDIO_SUFFIXES, *VIDEO_SUFFIXES}
AUDIO_UPLOAD_SUFFIXES = AUDIO_SUFFIXES
VIDEO_UPLOAD_SUFFIXES = VIDEO_SUFFIXES
UPLOAD_TYPE_TEXT = "文件类型暂不支持内容解析，但文件仍可以保存到工作区"
LEGACY_UPLOAD_TYPE_TEXT = "文件类型暂不支持，请上传 txt、md、pdf、docx、xlsx、csv、tsv、log、图片、音频或视频材料"
UPLOAD_BATCH_LIMIT = 5
UPLOAD_BATCH_PASSWORD = os.getenv("LAWRAG_UPLOAD_BATCH_PASSWORD", "")
PRIVATE_MATERIAL_HINTS = ("上传", "发你", "发给你", "材料", "资料", "文件", "附件", "图片", "照片", "视频", "录音", "聊天记录", "截图", "证据", "这份", "这个", "这些", "这段", "这张", "这条", "上述", "以下", "刚才", "刚刚", "刚上传", "刚发")
PRIVATE_FOLLOWUP_HINTS = ("帮我看", "帮我看看", "看一下", "看下", "看看", "分析一下", "分析下", "总结一下", "归纳一下", "识别一下", "提取一下", "里面", "上面", "前面", "内容", "能不能用", "能用吗", "有用吗", "有没有用", "可以用吗", "能不能作为证据", "能作为证据", "可以作为证据", "当证据", "能证明", "证明什么")
PUBLIC_LAW_LOOKUP_HINTS = ("民法典", "法条", "司法解释", "裁判规则", "判例", "案例", "法律规定", "法律依据")
PRIVATE_QUERY_WEAK_TERMS = {"根据", "以下", "上述", "这个", "那个", "这些", "那些", "可以", "是否", "当做", "作为", "已经", "我的", "意思", "什么", "怎么", "如何", "应该", "法律", "触犯", "解决", "上传", "材料", "资料", "文件", "附件", "图片", "照片", "视频", "录音", "截图", "证据", "这份", "这个", "这些", "这段", "这张", "这条", "刚才", "刚刚", "刚上传", "刚发"}
PRIVATE_TOKEN_PATTERN = re.compile(r"[\u4e00-\u9fffA-Za-z0-9]{2,}")
DOCUMENT_ID_PATTERN = re.compile(r"^[a-zA-Z0-9_-]+$")
RESET_CONTEXT_HINTS = ("和上次没关系", "和上次的没关系", "跟上次没关系", "与上次无关", "和之前没关系", "和之前的没关系", "跟之前没关系", "与之前无关", "不是上次", "不是之前", "不用之前", "别用之前", "不要用之前", "不要参考之前", "不要看之前", "不要看旧材料", "别看旧材料", "只看当前", "只认当前", "新问题", "新案子", "新事情", "另一个问题", "另一个案子", "另一个事情", "重新开始")
OLD_MATERIAL_REFERENCE_HINTS = ("之前上传", "上次上传", "以前上传", "历史上传", "旧材料", "旧文件", "旧资料", "之前发", "上次发", "以前发", "刚才上传", "刚刚上传", "刚上传", "刚才发", "刚刚发", "刚发", "上一轮", "上一个对话", "上次的材料", "之前的材料")
WINDOWS_RESERVED_NAMES = {"con", "prn", "aux", "nul", *(f"com{index}" for index in range(1, 10)), *(f"lpt{index}" for index in range(1, 10))}
INVALID_FILENAME_PATTERN = re.compile(r'[<>:"/\\|?*\x00-\x1f]')

def validate_document_id(document_id: str) -> bool:
    """检查文档或会话 ID 是否只包含允许的字母、数字、下划线和连字符。"""
    return bool(DOCUMENT_ID_PATTERN.match(str(document_id or "")))


def safe_upload_name(filename: str) -> str:
    """提取安全文件名，并拒绝路径、非法字符及 Windows 保留名称。"""
    name = Path(str(filename or "")).name.strip().rstrip(". ")
    if not name or name in {".", ".."} or INVALID_FILENAME_PATTERN.search(name) or Path(name).stem.lower() in WINDOWS_RESERVED_NAMES:
        raise ValueError("文件名不合法")
    return name


class LocalStore:
    """负责将用户上传文件安全地保存到工作区根目录并执行本地删除。"""

    def __init__(self, settings: Settings):
        """解析并创建配置指定的上传根目录。"""
        self.root = Path(settings.upload_dir).resolve()
        self.root.mkdir(parents=True, exist_ok=True)

    def save(self, user_id: str, filename: str, stream: BinaryIO, max_bytes: int | None = None) -> tuple[str, str, Path]:
        """分块写入上传流，限制文件大小，并返回文档 ID、安全名称和路径。"""
        document_id = uuid4().hex
        safe_name = safe_upload_name(filename)
        target = (self.root / user_id / document_id / safe_name).resolve()
        if self.root not in target.parents:
            raise ValueError("文件路径越界")
        target.parent.mkdir(parents=True, exist_ok=True)
        limit = max_bytes if max_bytes else 20 * 1024 * 1024
        total = 0
        try:
            with open(target, "wb") as handle:
                while True:
                    block = stream.read(64 * 1024)
                    if not block:
                        break
                    total += len(block)
                    if total > limit:
                        raise ValueError("文件大小超过限制")
                    handle.write(block)
        except Exception:
            target.unlink(missing_ok=True)
            raise
        return document_id, safe_name, target

    def delete(self, relative_path: str) -> None:
        """仅删除上传根目录内的指定文件，并尝试清理其空父目录。"""
        path = (self.root / relative_path).resolve()
        if self.root not in path.parents or not path.exists():
            return
        path.unlink(missing_ok=True)
        try:
            path.parent.rmdir()
        except OSError:
            pass


def read_image_with_rapidocr(path: Path) -> str:
    """调用 RapidOCR 识别图片并按检测顺序拼接非空文本。"""
    try:
        from rapidocr_onnxruntime import RapidOCR
    except ImportError as exc:
        raise ValueError("RapidOCR 未安装：请安装 rapidocr-onnxruntime 后重启服务") from exc
    result, _ = RapidOCR()(str(path))
    return "\n".join(str(item[1]).strip() for item in (result or []) if len(item) > 1 and str(item[1]).strip()).strip()


def _normalize_image_for_ocr(image):
    """修正图片方向、转换灰度，并放大小图以提高 OCR 识别率。"""
    from PIL import ImageOps

    image = ImageOps.exif_transpose(image).convert("L")
    if min(image.size) < 1200:
        image = image.resize((image.width * 2, image.height * 2))
    return image


def read_image_with_tesseract(path: Path, settings: Settings | None = None) -> str:
    """按配置调用 Tesseract 识别图片，并把常见环境错误转换为中文提示。"""
    try:
        from PIL import Image
        import pytesseract
        if settings and settings.tesseract_cmd:
            pytesseract.pytesseract.tesseract_cmd = settings.tesseract_cmd
        if settings and settings.tessdata_prefix:
            tessdata_path = Path(settings.tessdata_prefix).resolve()
            if tessdata_path.exists():
                os.environ["TESSDATA_PREFIX"] = str(tessdata_path)
        image = _normalize_image_for_ocr(Image.open(path))
        config = "--psm 6"
        return pytesseract.image_to_string(image, lang=settings.ocr_language if settings else "chi_sim+eng", config=config).strip()
    except Exception as exc:
        message = str(exc)
        if "tesseract is not installed" in message or "not in your PATH" in message:
            raise ValueError("Tesseract OCR 未启用：当前运行环境没有安装 Tesseract OCR，或 TESSERACT_CMD 配置不正确")
        if "Failed loading language" in message or "Error opening data file" in message:
            raise ValueError("Tesseract OCR 解析失败：当前服务器缺少 OCR 中文语言包 chi_sim，请安装 chi_sim.traineddata 或修正 TESSDATA_PREFIX")
        raise ValueError(f"Tesseract OCR 解析失败：{exc}")


def read_image_file(path: Path, settings: Settings | None = None) -> str:
    """按 OCR 配置选择 RapidOCR 或 Tesseract，并在 auto 模式下自动回退。"""
    if settings and not settings.enable_ocr:
        raise ValueError("图片 OCR 未启用：请将 ENABLE_OCR 设置为 true 后重启服务")
    engine = (settings.ocr_engine if settings else "auto").strip().lower() or "auto"
    errors: list[str] = []
    for name, reader in (("rapidocr", read_image_with_rapidocr), ("tesseract", read_image_with_tesseract)):
        if engine not in {"auto", name, "rapid" if name == "rapidocr" else name}:
            continue
        try:
            text = reader(path) if name == "rapidocr" else reader(path, settings)
            if text.strip() or engine != "auto":
                return text.strip()
        except ValueError as exc:
            errors.append(str(exc))
            if engine != "auto":
                raise
    if errors:
        raise ValueError("图片 OCR 解析失败：" + "；".join(errors))
    raise ValueError(f"图片 OCR 引擎不支持：{engine}")


def read_text_file(path: Path) -> str:
    """以 UTF-8 容错模式读取纯文本文件。"""
    return path.read_text(encoding="utf-8", errors="ignore")


def read_pdf_file(path: Path) -> str:
    """逐页提取 PDF 文本，并附加页码标记以便后续引用定位。"""
    from pypdf import PdfReader
    pages = []
    for index, page in enumerate(PdfReader(str(path)).pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"【PDF第{index}页】\n{text}")
    return "\n\n".join(pages)


def pdf_has_tables(path: Path, max_pages: int = 5) -> bool:
    """只检查前几页是否存在表格，让普通 PDF 保持快速解析。"""
    try:
        import pdfplumber

        with pdfplumber.open(path) as document:
            return any(page.find_tables() for page in document.pages[:max_pages])
    except Exception:
        return False


def read_pdf_with_ocr(path: Path, settings: Settings | None = None) -> str:
    """把 PDF 页面渲染为图片后逐页 OCR，跳过无法识别的页面。"""
    try:
        from pdf2image import convert_from_path
    except ImportError:
        return ""
    try:
        images = convert_from_path(str(path), dpi=220, poppler_path=str(getattr(settings, "poppler_path", "") or "") or None)
    except Exception as exc:
        logger.warning("PDF OCR 渲染失败", extra={"event": "pdf_ocr_render_failed", "fields": {"file_name": path.name, "error": str(exc)}})
        return ""
    if not images:
        return ""
    parts: list[str] = []
    for index, image in enumerate(images, start=1):
        try:
            normalized = _normalize_image_for_ocr(image)
            with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as handle:
                temp_path = Path(handle.name)
            try:
                normalized.save(temp_path)
                page_text = read_image_file(temp_path, settings).strip()
            finally:
                temp_path.unlink(missing_ok=True)
        except Exception as exc:
            logger.warning("PDF 单页 OCR 失败", extra={"event": "pdf_ocr_page_failed", "fields": {"file_name": path.name, "page": index, "error": str(exc)}})
            continue
        if page_text:
            parts.append(f"【PDF OCR第{index}页】\n{page_text}")
    return "\n\n".join(parts)


def read_docx_file(path: Path) -> str:
    """提取 Word 文档中的段落和表格文本，同时保留表格序号。"""
    from docx import Document
    document = Document(str(path))
    lines = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table_index, table in enumerate(document.tables, start=1):
        lines.append(f"【Word表格{table_index}】")
        for row in table.rows:
            values = [cell.text.strip().replace("\n", " ") for cell in row.cells if cell.text.strip()]
            if values:
                lines.append(" | ".join(values))
    return "\n".join(lines)


def read_xlsx_file(path: Path) -> str:
    """以只读模式提取 Excel 各工作表的非空单元格并标记行号。"""
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    lines: list[str] = []
    try:
        for sheet in workbook.worksheets:
            lines.append(f"【工作表】{sheet.title}")
            for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                values = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
                if values:
                    lines.append(f"第{row_index}行：" + " | ".join(values))
    finally:
        workbook.close()
    return "\n".join(lines)


def parse_text(path: Path, settings: Settings | None = None) -> str:
    """根据文件扩展名调度对应文本、PDF、Office 或图片解析器。"""
    suffix = path.suffix.lower()
    if suffix in TEXT_SUFFIXES:
        return read_text_file(path)
    if suffix == ".pdf":
        return read_pdf_file(path)
    if suffix == ".docx":
        return read_docx_file(path)
    if suffix == ".xlsx":
        return read_xlsx_file(path)
    if suffix in IMAGE_SUFFIXES:
        return read_image_file(path, settings)
    return ""


def media_type(path: Path) -> str:
    """根据扩展名把文件归类为 image、audio、video、text 或 file。"""
    suffix = path.suffix.lower()
    if suffix in IMAGE_SUFFIXES:
        return "image"
    if suffix in AUDIO_SUFFIXES:
        return "audio"
    if suffix in VIDEO_SUFFIXES:
        return "video"
    return "text" if suffix in ALLOWED_UPLOAD_SUFFIXES else "file"


def extraction_metadata(path: Path, text: str, method: str | None = None) -> dict:
    """根据文件类型和抽取方式生成统一的 OCR、多模态及预览元数据。"""
    kind = media_type(path)
    multimodal = method in {"multimodal", "ocr_multimodal"}
    ocr = method in {"ocr", "ocr_multimodal", "ocr_pdf", "ocr_pdf_multimodal"} or (method is None and kind == "image")
    mineru = method == "mineru"
    return {
        "document_type": kind,
        "media_type": kind,
        "extracted_char_count": len(text.strip()),
        "extraction_method": method or ("ocr" if kind == "image" else "text_parser"),
        "ocr_used": ocr,
        "ocr_status": "success" if ocr else "not_used",
        "multimodal_used": multimodal,
        "multimodal_status": "success" if multimodal else "not_used",
        "mineru_used": mineru,
        "mineru_status": "success" if mineru else "not_used",
        "extraction_preview": " ".join(text.strip().split())[:1000],
    }


def merge_image_text(ocr_text: str, multimodal_text: str) -> str:
    """合并图片 OCR 与多模态分析文本，并避免重复包含相同内容。"""
    ocr_value = str(ocr_text or "").strip()
    multimodal_value = str(multimodal_text or "").strip()
    if not ocr_value:
        return multimodal_value
    if not multimodal_value or ocr_value in multimodal_value:
        return multimodal_value or ocr_value
    return f"【OCR识别】\n{ocr_value}\n\n【多模态分析】\n{multimodal_value}"


def _extract_text_without_mineru(path: Path, safe_name: str, settings: Settings, model: ModelGateway) -> tuple[str, dict]:
    """按媒体类型编排文本解析、OCR 和多模态模型，并返回文本及元数据。"""
    kind = media_type(path)
    if kind in {"audio", "video"}:
        text = model.analyze_media(path, kind, safe_name)
        return text, extraction_metadata(path, text, "multimodal")
    if kind == "image":
        ocr_text = ""
        try:
            ocr_text = parse_text(path, settings)
        except Exception:
            logger.warning("图片 OCR 解析失败，回退多模态分析", extra={"event": "ocr_image_fallback_to_multimodal", "fields": {"file_name": safe_name}})
        try:
            text = model.analyze_media(path, kind, safe_name)
            if text.strip():
                merged = merge_image_text(ocr_text, text)
                return merged, extraction_metadata(path, merged, "ocr_multimodal" if ocr_text.strip() else "multimodal")
        except Exception:
            logger.warning("图片多模态分析失败，保留 OCR 识别结果", extra={"event": "multimodal_image_fallback_to_ocr", "fields": {"file_name": safe_name}})
        if ocr_text.strip():
            return ocr_text, extraction_metadata(path, ocr_text, "ocr")
        raise ValueError("图片内容无法识别：OCR 和多模态分析都未返回可检索文本")
    if path.suffix.lower() == ".pdf":
        text = parse_text(path, settings)
        if not text.strip() or len(text.strip()) < max(80, int(getattr(settings, "ocr_min_text_chars", 80) or 80)):
            ocr_text = read_pdf_with_ocr(path, settings)
            if text.strip() and ocr_text.strip():
                merged = f"【PDF文本抽取】\n{text.strip()}\n\n【PDF OCR识别】\n{ocr_text.strip()}"
                return merged, extraction_metadata(path, merged, "ocr_pdf_multimodal")
            if ocr_text.strip():
                return ocr_text, extraction_metadata(path, ocr_text, "ocr_pdf")
        return text, extraction_metadata(path, text, "text_parser")
    text = parse_text(path, settings)
    return text, extraction_metadata(path, text, "text_parser")


def extract_text(path: Path, safe_name: str, settings: Settings, model: ModelGateway) -> tuple[str, dict]:
    """简单文件本地解析；图片、扫描件和表格 PDF 优先使用 MinerU。"""
    kind = media_type(path)
    mineru_ready = bool(settings.mineru_enabled and settings.mineru_api_token)
    if not mineru_ready or (kind != "image" and path.suffix.lower() != ".pdf"):
        return _extract_text_without_mineru(path, safe_name, settings, model)

    force_ocr = kind == "image"
    if path.suffix.lower() == ".pdf":
        local_text = parse_text(path, settings)
        minimum_text = max(80, int(settings.ocr_min_text_chars or 80))
        force_ocr = len(local_text.strip()) < minimum_text
        if not force_ocr and not pdf_has_tables(path):
            return local_text, extraction_metadata(path, local_text, "text_parser")

    try:
        text = extract_with_mineru(path, settings, force_ocr=force_ocr)
        return text, extraction_metadata(path, text, "mineru")
    except Exception as exc:
        logger.warning(
            "MinerU 解析失败，回退到原有解析流程",
            extra={
                "event": "mineru_fallback",
                "fields": {"file_name": safe_name, "error_type": type(exc).__name__},
            },
        )
        text, metadata = _extract_text_without_mineru(path, safe_name, settings, model)
        metadata["mineru_used"] = True
        metadata["mineru_status"] = "failed"
        return text, metadata
