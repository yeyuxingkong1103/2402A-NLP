import json
from hashlib import sha256
from pathlib import Path

from docx import Document

from data_pipeline.load import file_sha256, load_docx, load_json, load_text
from data_pipeline.clean import clean_public_collections, clean_text
from data_pipeline.chunk import chunk_records
from data_pipeline.parse import (
    build_citations,
    extract_articles,
    extract_cases,
    extract_elements,
    extract_evidence,
    extract_interpretations,
    extract_processes,
    extract_questions,
)
from data_pipeline.public_main import process_public_data


def test_loaders_read_the_three_public_source_formats(tmp_path) -> None:
    text_path = tmp_path / "sample.txt"
    text_path.write_text("第一条 示例内容", encoding="utf-8")

    json_path = tmp_path / "sample.json"
    json_path.write_text(json.dumps({"items": [1]}, ensure_ascii=False), encoding="utf-8")

    docx_path = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("司法解释标题")
    document.add_paragraph("第一条 解释内容")
    document.save(docx_path)

    assert load_text(text_path) == "第一条 示例内容"
    assert load_json(json_path) == {"items": [1]}
    assert load_docx(docx_path) == "司法解释标题\n第一条 解释内容"


def test_file_hash_is_stable_and_changes_with_content(tmp_path) -> None:
    path = tmp_path / "source.txt"
    path.write_text("内容一", encoding="utf-8")
    first = file_sha256(path)

    assert first == file_sha256(path)

    path.write_text("内容二", encoding="utf-8")
    assert file_sha256(path) != first


def test_law_articles_keep_article_number_and_hierarchy() -> None:
    rows = extract_articles(
        "第一编 总则\n第一章 基本规定\n第一条 第一条内容。\n第二条 第二条内容。"
    )

    assert [row["article_number"] for row in rows] == [1, 2]
    assert rows[0]["chapter"] == "第一编 总则 > 第一章 基本规定"
    assert rows[0]["id"] == "civil_code_articles_article_1"


def test_interpretations_are_split_by_document_and_article() -> None:
    text = """最高人民法院关于合同纠纷的解释
法释〔2025〕1号
第一条 根据民法典第一条制定本条。
第二条 第二条解释。
最高人民法院关于婚姻家庭的解释
法释〔2025〕2号
第一条 根据民法典第一千零七十九条处理。
"""

    rows = extract_interpretations(text)

    assert len(rows) == 3
    assert rows[0]["law_name"] == "最高人民法院关于合同纠纷的解释"
    assert rows[0]["related_article_numbers"] == [1]
    assert rows[2]["related_article_numbers"] == [1079]


def test_cases_are_structured_and_non_case_material_is_filtered() -> None:
    text = """【案例 1】
标题：买卖合同纠纷案
案号：（2025）测试1号
裁判日期：2025-01-01
审理法院：测试法院
案由：买卖合同纠纷
文书种类：民事判决书
【案情内容】
法院依据民法典第五百七十七条判决被告承担违约责任。
【案例 2】
标题：征求意见稿
文书种类：司法解释征求意见稿
材料类型：司法解释/意见稿
【案情内容】
这不是案例。
"""

    rows = extract_cases(text)

    assert len(rows) == 1
    assert rows[0]["title"] == "买卖合同纠纷案"
    assert rows[0]["related_article_numbers"] == [577]


def test_evidence_process_question_and_elements_are_structured() -> None:
    evidence = extract_evidence(
        "案例编号：0001\n源案例编号：case-1\n详细内容：\n规则名称：电子证据规则\n"
        "规则摘要：应提交原始载体。\n证据规则：\n依据民法典第一条。"
    )
    processes = extract_processes(
        "# CPC-0001 合同纠纷\n适用阶段：立案、审理\n适用场景：\n处理合同纠纷。\n"
        "处理路径：\n1. 收集合同。\n2. 整理时间轴。\n所需材料：\n- 合同\n"
        "交付成果：\n- 时间轴\n来源文件：\n- 民法典\n依据数量：1\n法律依据：\n民法典第一条"
    )
    questions = extract_questions(
        "# 1\n问题：\n合同违约怎么办？\n回答：\n可以要求承担违约责任。\n"
        "来源：测试\ncomplexity：2\nclarity：3\ninformativeness：4"
    )
    elements = extract_elements(
        {"civil_elements": [{"serial_number": 1, "title": "第一条要件", "case_summary": "要件内容"}]}
    )

    assert evidence[0]["evidence_id"] == "civil_evidence_0001"
    assert processes[0]["steps"] == ["收集合同。", "整理时间轴。"]
    assert questions[0]["question"] == "合同违约怎么办？"
    assert elements[0]["serial_number"] == "1"


def test_citations_are_derived_from_article_references() -> None:
    collections = {
        "civil_code_articles": [
            {"id": "civil_code_articles_article_1", "article_number": 1}
        ],
        "civil_cases": [
            {
                "case_id": "case-1",
                "title": "案例",
                "related_article_numbers": [1],
            }
        ],
    }

    citations = build_citations(collections)

    assert len(citations) == 1
    assert citations[0]["source_id"] == "case-1"
    assert citations[0]["target_id"] == "civil_code_articles_article_1"


def test_cleaning_removes_noise_duplicates_and_placeholder_cases() -> None:
    collections = {
        "civil_code_articles": [
            {"id": "article-1", "article_content": " 第一条\t内容。\n-------- "},
            {"id": "article-1", "article_content": "重复内容"},
        ],
        "civil_cases": [
            {"case_id": "case-1", "title": "有效案例", "summary": "有效案情"},
            {"case_id": "case-2", "title": "占位案例", "summary": "案情内容暂缺"},
        ],
    }

    cleaned, report = clean_public_collections(collections)

    assert cleaned["civil_code_articles"] == [
        {"id": "article-1", "article_content": "第一条 内容。"}
    ]
    assert [row["case_id"] for row in cleaned["civil_cases"]] == ["case-1"]
    assert report["civil_code_articles"]["duplicate_records"] == 1
    assert report["civil_cases"]["placeholder_records"] == 1


def test_clean_text_preserves_paragraphs_but_removes_page_markers() -> None:
    assert clean_text("第一段\r\n— 12 —\r\n\r\n第二段") == "第一段\n第二段"


def test_chunking_keeps_short_records_and_splits_long_records_stably() -> None:
    rows = [
        {
            "case_id": "case-1",
            "title": "测试案例",
            "summary": "第一句内容。第二句内容。第三句内容。第四句内容。",
        }
    ]

    first = chunk_records("civil_cases", rows, chunk_size=12, overlap=3)
    second = chunk_records("civil_cases", rows, chunk_size=12, overlap=3)

    assert len(first) > 1
    assert first == second
    assert all(item["record_id"] == "case-1" for item in first)
    assert all(len(item["content"]) <= 12 for item in first)
    assert [item["chunk_index"] for item in first] == list(range(len(first)))


def _write_sample_public_data(data_dir: Path) -> None:
    sources = {
        "civil_code_articles/civil_code_articles.txt": "第一编 总则\n第一章 基本规定\n第一条 法条内容。",
        "civil_cases/civil_cases.txt": (
            "【案例 1】\n标题：合同案例\n案号：测试1号\n文书种类：民事判决书\n"
            "【案情内容】\n依据民法典第一条作出判决。\n"
            "【案例 2】\n标题：合同案例\n案号：测试1号\n文书种类：民事判决书\n"
            "【案情内容】\n依据民法典第一条作出判决。\n"
            "【案例 3】\n标题：征求意见稿\n文书种类：司法解释征求意见稿\n"
            "材料类型：司法解释/意见稿\n【案情内容】\n这不是民事案例。"
        ),
        "civil_evidence/civil_evidence.txt": (
            "案例编号：0001\n源案例编号：case-1\n详细内容：\n规则名称：证据规则\n"
            "规则摘要：保留原件。\n依据民法典第一条。"
        ),
        "civil_processes/civil_processes.txt": (
            "# CPC-0001 合同流程\n适用阶段：立案\n适用场景：\n处理纠纷。\n处理路径：\n"
            "1. 收集材料。\n所需材料：\n- 合同\n交付成果：\n- 清单\n来源文件：\n"
            "- 民法典\n依据数量：1\n法律依据：\n民法典第一条"
        ),
        "civil_questions/civil_questions.txt": (
            "# 1\n问题：\n如何处理？\n回答：\n依据民法典第一条处理。\n来源：测试\n"
            "complexity：1\nclarity：1\ninformativeness：1"
        ),
    }
    for relative_path, content in sources.items():
        path = data_dir / relative_path
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(content, encoding="utf-8")

    elements_path = data_dir / "civil_elements/civil_elements.json"
    elements_path.parent.mkdir(parents=True, exist_ok=True)
    elements_path.write_text(
        json.dumps(
            {"civil_elements": [{"serial_number": 1, "title": "第一条要件", "case_summary": "要件内容"}]},
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    interpretations_path = data_dir / "civil_interpretations/civil_interpretations.docx"
    interpretations_path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_paragraph("最高人民法院关于测试事项的解释")
    document.add_paragraph("法释〔2025〕1号")
    document.add_paragraph("第一条 根据民法典第一条处理。")
    document.save(interpretations_path)


def test_public_pipeline_writes_eight_collections_and_reports(tmp_path) -> None:
    data_dir = tmp_path / "public"
    output_dir = tmp_path / "processed"
    _write_sample_public_data(data_dir)

    result = process_public_data(data_dir, output_dir)

    expected = {
        "civil_code_articles",
        "civil_interpretations",
        "civil_cases",
        "civil_elements",
        "civil_evidence",
        "civil_processes",
        "civil_questions",
        "civil_citations",
    }
    assert set(result["record_counts"]) == expected
    assert all((output_dir / f"{name}.json").is_file() for name in expected)
    assert all((output_dir / "chunks" / f"{name}_chunks.json").is_file() for name in expected)
    assert (output_dir / "quality_report.json").is_file()
    assert (output_dir / "build_manifest.json").is_file()
    assert "embedding" not in (output_dir / "civil_cases.json").read_text(encoding="utf-8")
    assert result["quality_report"]["civil_cases"]["input_records"] == 3
    assert result["quality_report"]["civil_cases"]["parsed_records"] == 1
    assert result["quality_report"]["civil_cases"]["parser_dropped_records"] == 2
    assert result["quality_report"]["civil_cases"]["filter_reasons"] == {
        "duplicate_record": 1,
        "judicial_interpretation_material": 1,
    }
    assert set(result["chunk_counts"]) == expected


def test_public_pipeline_collection_output_is_stable(tmp_path) -> None:
    data_dir = tmp_path / "public"
    output_dir = tmp_path / "processed"
    _write_sample_public_data(data_dir)

    process_public_data(data_dir, output_dir)
    first = sha256((output_dir / "civil_citations.json").read_bytes()).hexdigest()
    process_public_data(data_dir, output_dir)
    second = sha256((output_dir / "civil_citations.json").read_bytes()).hexdigest()

    assert first == second
