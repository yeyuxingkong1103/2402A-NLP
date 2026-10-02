"""读取模块测试：只用临时示例文件，不读取公共资料、不连接数据库。"""
import importlib
import json
import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, NumberObject, DecodedStreamObject

reader = importlib.import_module("data_pipeline.load")


def make_pdf(path, pages):
    """生成真实PDF；每页可以有文字，也可以有一张图片。"""
    writer = PdfWriter()
    for text, has_image in pages:
        page = writer.add_blank_page(width=400, height=400)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        resources = DictionaryObject({NameObject("/Font"): DictionaryObject(
            {NameObject("/F1"): writer._add_object(font)})})
        content = f"BT /F1 12 Tf 30 350 Td ({text}) Tj ET\n" if text else ""
        if has_image:
            image = DecodedStreamObject()
            image.set_data(bytes([255, 255, 255]) * 100)
            image.update({NameObject("/Type"): NameObject("/XObject"),
                          NameObject("/Subtype"): NameObject("/Image"),
                          NameObject("/Width"): NumberObject(10),
                          NameObject("/Height"): NumberObject(10),
                          NameObject("/ColorSpace"): NameObject("/DeviceRGB"),
                          NameObject("/BitsPerComponent"): NumberObject(8)})
            resources[NameObject("/XObject")] = DictionaryObject(
                {NameObject("/Im1"): writer._add_object(image)})
            content += "q 200 0 0 200 20 20 cm /Im1 Do Q\n"
        page[NameObject("/Resources")] = resources
        stream = DecodedStreamObject()
        stream.set_data(content.encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(path)


def test_text_removes_bom_without_losing_chinese(tmp_path):
    path = tmp_path / "law.txt"
    path.write_text("第一条 内容", encoding="utf-8-sig")
    assert reader.load_text(path) == "第一条 内容"


def test_wrong_encoding_is_not_silently_ignored(tmp_path):
    path = tmp_path / "law.txt"
    path.write_bytes("法律条文".encode("gb18030"))
    with pytest.raises(UnicodeError):
        reader.load_text(path)


def test_text_can_use_explicit_encoding(tmp_path):
    path = tmp_path / "law.txt"
    path.write_bytes("法律条文".encode("gb18030"))
    assert reader.load_text(path, encoding="gb18030") == "法律条文"


def test_word_keeps_table_between_its_paragraphs(tmp_path):
    path = tmp_path / "law.docx"
    document = Document()
    document.add_paragraph("表格前")
    table = document.add_table(rows=1, cols=3)
    table.cell(0, 0).text = "项目"
    table.cell(0, 2).text = "金额"
    document.add_paragraph("表格后")
    document.save(path)
    text = reader.load_docx(path)
    assert text.index("表格前") < text.index("项目") < text.index("表格后")
    assert "项目 |  | 金额" in text


def test_excel_keeps_cell_addresses_and_formula(tmp_path):
    path = tmp_path / "amount.xlsx"
    workbook = Workbook()
    workbook.active.title = "清单"
    workbook.active.append(["金额", None, 1000])
    workbook.active["C2"] = "=C1*2"
    workbook.save(path)
    text = reader.load_xlsx(path)
    assert "A1=金额 | B1= | C1=1000" in text
    assert "C2==C1*2" in text


def test_mixed_pdf_reads_scanned_page_even_when_other_pages_have_text(tmp_path, monkeypatch):
    path = tmp_path / "mixed.pdf"
    make_pdf(path, [("A long legal text. " * 8, False), ("", True)])
    monkeypatch.setattr(reader, "render_pdf_page", lambda path, page: Image.new("RGB", (100, 100)), raising=False)
    monkeypatch.setattr(reader, "ocr_image", lambda image: ("第二页扫描文字", "rapidocr"), raising=False)
    result = reader.extract_content(path)
    assert "第二页扫描文字" in result.text
    assert [page["page"] for page in result.pages] == [1, 2]
    assert result.ocr_used is True


def test_text_pdf_preserves_real_pages_without_ocr(tmp_path, monkeypatch):
    path = tmp_path / "text.pdf"
    make_pdf(path, [("Article one", False), ("Article two", False)])
    def unexpected_ocr(*args):
        pytest.fail("纯文字页面不应该调用OCR")
    monkeypatch.setattr(reader, "ocr_image", unexpected_ocr, raising=False)
    result = reader.extract_content(path)
    assert [page["page"] for page in result.pages] == [1, 2]
    assert result.pages[1]["text"] == "Article two"
    assert result.ocr_used is False


def test_blank_pdf_page_is_not_an_ocr_failure(tmp_path):
    path = tmp_path / "blank.pdf"
    make_pdf(path, [("Article one", False), ("", False)])
    result = reader.extract_content(path)
    assert len(result.pages) == 2
    assert result.pages[1]["text"] == ""
    assert result.status == "success"


def test_failed_pdf_page_is_reported_and_string_loader_refuses_partial_result(tmp_path, monkeypatch):
    path = tmp_path / "mixed.pdf"
    make_pdf(path, [("Article one", False), ("", True)])
    def failed_render(*args):
        raise RuntimeError("测试：缺少渲染工具")
    monkeypatch.setattr(reader, "render_pdf_page", failed_render, raising=False)
    result = reader.extract_content(path)
    assert result.status == "partial"
    assert result.pages[1]["status"] == "failed"
    assert any("第2页" in warning for warning in result.warnings)
    with pytest.raises(RuntimeError, match="完整"):
        reader.load_pdf(path)


def test_image_uses_rapidocr_first(tmp_path, monkeypatch):
    path = tmp_path / "scan.png"
    Image.new("RGB", (100, 100)).save(path)
    monkeypatch.setattr(reader, "read_image_with_rapidocr", lambda image: "Rapid识别文字", raising=False)
    def unexpected_tesseract(*args):
        pytest.fail("RapidOCR成功后不应再调用备用引擎")
    monkeypatch.setattr(reader, "read_image_with_tesseract", unexpected_tesseract, raising=False)
    assert reader.load_image(path) == "Rapid识别文字"


def test_image_falls_back_to_tesseract(tmp_path, monkeypatch):
    path = tmp_path / "scan.png"
    Image.new("RGB", (100, 100)).save(path)
    def missing_rapid(*args):
        raise ImportError("未安装")
    monkeypatch.setattr(reader, "read_image_with_rapidocr", missing_rapid, raising=False)
    monkeypatch.setattr(reader, "read_image_with_tesseract", lambda image: "备用识别文字", raising=False)
    result = reader.extract_content(path)
    assert result.text == "备用识别文字"
    assert result.extract_method == "tesseract"


def test_empty_ocr_result_is_not_success(tmp_path, monkeypatch):
    path = tmp_path / "scan.png"
    Image.new("RGB", (100, 100)).save(path)
    monkeypatch.setattr(reader, "read_image_with_rapidocr", lambda image: "", raising=False)
    monkeypatch.setattr(reader, "read_image_with_tesseract", lambda image: "", raising=False)
    with pytest.raises(RuntimeError, match="OCR"):
        reader.load_image(path)


def test_office_document_does_not_invent_a_page_number(tmp_path):
    path = tmp_path / "law.docx"
    document = Document()
    document.add_paragraph("条文")
    document.save(path)
    result = reader.extract_content(path)
    assert result.pages[0]["page"] is None


def test_json_keeps_structure_and_utf8_bom(tmp_path):
    path = tmp_path / "law.json"
    path.write_text('{"标题": "条文", "编号": 1}', encoding="utf-8-sig")
    assert reader.load_json(path) == {"标题": "条文", "编号": 1}
    assert json.loads(reader.load(path)) == {"标题": "条文", "编号": 1}


def test_wrong_file_format_reports_error(tmp_path):
    path = tmp_path / "wrong.pdf"
    path.write_text("这不是PDF", encoding="utf-8")
    with pytest.raises(Exception):
        reader.load(path)


def test_unsupported_suffix_reports_error(tmp_path):
    path = tmp_path / "old.doc"
    path.write_bytes(b"not supported")
    with pytest.raises(ValueError, match="不支持"):
        reader.load(path)


def test_pdf_page_renders_without_poppler(tmp_path, monkeypatch):
    path = tmp_path / "scan.pdf"
    make_pdf(path, [("", True)])
    monkeypatch.delenv("POPPLER_PATH", raising=False)
    monkeypatch.setenv("PATH", "")
    image = reader.render_pdf_page(path, 1)
    try:
        assert image.width > 400
        assert image.height > 400
    finally:
        image.close()


def test_pdf_table_preserves_rows_and_does_not_repeat_cell_text(tmp_path):
    path = tmp_path / "table.pdf"
    make_pdf(path, [("", False)])
    writer = PdfWriter(clone_from=path)
    stream = DecodedStreamObject()
    stream.set_data(b"""BT /F1 12 Tf 30 350 Td (Before) Tj ET
30 220 m 230 220 l S 30 260 m 230 260 l S 30 300 m 230 300 l S
30 220 m 30 300 l S 130 220 m 130 300 l S 230 220 m 230 300 l S
BT /F1 12 Tf 40 280 Td (Item) Tj ET
BT /F1 12 Tf 140 280 Td (Amount) Tj ET
BT /F1 12 Tf 40 240 Td (Rent) Tj ET
BT /F1 12 Tf 140 240 Td (1000) Tj ET
BT /F1 12 Tf 30 180 Td (After) Tj ET
""")
    writer.pages[0][NameObject("/Contents")] = writer._add_object(stream)
    writer.write(path)
    result = reader.extract_content(path)
    assert result.pages[0]["tables"][0]["rows"] == [["Item", "Amount"], ["Rent", "1000"]]
    assert result.text.count("1000") == 1
    assert result.text.index("Before") < result.text.index("Item") < result.text.index("After")


def test_file_size_limit_is_checked_before_reading(tmp_path, monkeypatch):
    path = tmp_path / "big.txt"
    path.write_text("too big", encoding="utf-8")
    monkeypatch.setattr(reader, "MAX_FILE_BYTES", 4)
    with pytest.raises(ValueError, match="拆分"):
        reader.load(path)


def test_pdf_page_limit_is_checked_before_processing(tmp_path, monkeypatch):
    path = tmp_path / "big.pdf"
    make_pdf(path, [("one", False), ("two", False)])
    monkeypatch.setattr(reader, "MAX_PDF_PAGES", 1)
    with pytest.raises(ValueError, match="拆分"):
        reader.extract_content(path)


def test_real_rapidocr_reads_generated_image(tmp_path):
    # 真正运行本地ONNX模型，不下载模型，也不调用外部接口。
    path = tmp_path / "ocr.png"
    image = Image.new("RGB", (1000, 240), "white")
    ImageDraw.Draw(image).text((50, 70), "LEGAL 12345", font=ImageFont.load_default(size=64), fill="black")
    image.save(path)
    result = reader.extract_content(path)
    assert result.extract_method == "rapidocr"
    assert "LEGAL" in result.text.upper()
    assert "12345" in result.text


def test_real_scanned_pdf_runs_rendering_and_ocr(tmp_path):
    path = tmp_path / "scan.pdf"
    image = Image.new("RGB", (1000, 240), "white")
    ImageDraw.Draw(image).text((50, 70), "LEGAL 12345", font=ImageFont.load_default(size=64), fill="black")
    writer = PdfWriter()
    page = writer.add_blank_page(width=600, height=300)
    picture = DecodedStreamObject()
    picture.set_data(image.tobytes())
    picture.update({NameObject("/Type"): NameObject("/XObject"), NameObject("/Subtype"): NameObject("/Image"),
                    NameObject("/Width"): NumberObject(1000), NameObject("/Height"): NumberObject(240),
                    NameObject("/ColorSpace"): NameObject("/DeviceRGB"), NameObject("/BitsPerComponent"): NumberObject(8)})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/XObject"): DictionaryObject(
        {NameObject("/Scan"): writer._add_object(picture)})})
    stream = DecodedStreamObject()
    stream.set_data(b"q 500 0 0 120 50 100 cm /Scan Do Q")
    page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(path)
    result = reader.extract_content(path)
    assert result.status == "success"
    assert result.ocr_used is True
    assert result.pages[0]["method"] == "rapidocr"
    assert "12345" in result.text


def test_pdf_drawing_without_text_layer_attempts_ocr(tmp_path, monkeypatch):
    path = tmp_path / "drawing.pdf"
    make_pdf(path, [("", False)])
    writer = PdfWriter(clone_from=path)
    stream = DecodedStreamObject()
    stream.set_data(b"20 20 100 100 re f")
    writer.pages[0][NameObject("/Contents")] = writer._add_object(stream)
    writer.write(path)
    monkeypatch.setattr(reader, "render_pdf_page", lambda path, page: Image.new("RGB", (100, 100)))
    monkeypatch.setattr(reader, "ocr_image", lambda image: ("图形文字", "rapidocr"))
    result = reader.extract_content(path)
    assert "图形文字" in result.text


def test_text_and_image_on_same_page_requires_review_instead_of_silent_success(tmp_path):
    path = tmp_path / "mixed_page.pdf"
    make_pdf(path, [("A long legal text. " * 8, True)])
    result = reader.extract_content(path)
    assert result.status == "partial"
    assert result.pages[0]["status"] == "needs_review"
    assert result.warnings
    with pytest.raises(RuntimeError, match="完整"):
        reader.load_pdf(path)


def test_multiframe_tiff_reads_every_frame(tmp_path, monkeypatch):
    path = tmp_path / "pages.tiff"
    first = Image.new("RGB", (100, 100), "red")
    second = Image.new("RGB", (100, 100), "blue")
    first.save(path, save_all=True, append_images=[second])
    def frame_ocr(image):
        color = image.convert("RGB").getpixel((0, 0))
        return ("第一帧" if color[0] > color[2] else "第二帧"), "rapidocr"
    monkeypatch.setattr(reader, "ocr_image", frame_ocr)
    result = reader.extract_content(path)
    assert len(result.pages) == 2
    assert result.pages[1]["text"] == "第二帧"
    assert "第一帧" in result.text and "第二帧" in result.text


def test_pdf_with_only_graphics_state_commands_is_blank_not_ocr_failure(tmp_path, monkeypatch):
    path = tmp_path / "blank_state.pdf"
    make_pdf(path, [("Article one", False), ("", False)])
    writer = PdfWriter(clone_from=path)
    stream = DecodedStreamObject()
    stream.set_data(b"q Q")
    writer.pages[1][NameObject("/Contents")] = writer._add_object(stream)
    writer.write(path)
    def unexpected_ocr(image):
        pytest.fail("只有状态操作的空白页不应该尝试OCR")
    monkeypatch.setattr(reader, "ocr_image", unexpected_ocr)
    result = reader.extract_content(path)
    assert result.status == "success"
    assert result.pages[1]["text"] == ""


def test_failed_tiff_frame_is_not_silently_skipped(tmp_path, monkeypatch):
    path = tmp_path / "pages.tiff"
    first = Image.new("RGB", (100, 100), "red")
    second = Image.new("RGB", (100, 100), "blue")
    first.save(path, save_all=True, append_images=[second])
    def frame_ocr(image):
        color = image.convert("RGB").getpixel((0, 0))
        if color[0] < color[2]:
            raise RuntimeError("第二帧失败")
        return "第一帧", "rapidocr"
    monkeypatch.setattr(reader, "ocr_image", frame_ocr)
    result = reader.extract_content(path)
    assert result.status == "partial"
    assert result.pages[1]["status"] == "failed"
    assert any("第2帧" in warning for warning in result.warnings)
    with pytest.raises(RuntimeError, match="完整"):
        reader.load_image(path)
