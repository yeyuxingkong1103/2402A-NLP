from __future__ import annotations

from dataclasses import asdict, dataclass
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile


TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".rst", ".csv", ".tsv", ".log", ".json"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}
SUPPORTED_SUFFIXES = TEXT_SUFFIXES | IMAGE_SUFFIXES | {".pdf", ".docx", ".xlsx"}


@dataclass
class ExtractedDocument:
    text: str
    pages: list[dict]
    extract_method: str
    ocr_used: bool
    vision_used: bool

    def to_dict(self) -> dict:
        return asdict(self)


def load_text(path: Path) -> str:
    return Path(path).read_text(encoding="utf-8", errors="ignore")


def load_json(path: Path) -> str:
    return Path(path).read_text(encoding="utf-8-sig", errors="ignore")


def load_pdf(path: Path) -> str:
    from pypdf import PdfReader

    from PIL import Image, ImageOps
    from pdf2image import convert_from_path
    import pytesseract

    def normalize(image):
        image = ImageOps.exif_transpose(image).convert("L")
        if min(image.size) < 1200:
            image = image.resize((image.width * 2, image.height * 2))
        return image

    pages = []
    text_pages = []
    for index, page in enumerate(PdfReader(str(path)).pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            text_pages.append(f"【PDF第{index}页】\n{text}")
    text_output = "\n\n".join(text_pages)
    if len(text_output.strip()) >= 80:
        return text_output

    try:
        images = convert_from_path(str(path), dpi=220)
    except Exception:
        return text_output

    for index, image in enumerate(images, start=1):
        try:
            normalized = normalize(image)
            page_text = pytesseract.image_to_string(normalized, lang="chi_sim+eng", config="--psm 6").strip()
        except Exception:
            continue
        if page_text:
            pages.append(f"【PDF OCR第{index}页】\n{page_text}")

    ocr_text = "\n\n".join(pages)
    if text_output and ocr_text:
        return f"【PDF文本抽取】\n{text_output}\n\n【PDF OCR识别】\n{ocr_text}"
    return ocr_text or text_output


def _load_docx_xml(path: Path) -> str:
    with ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))
    paragraphs = []
    for paragraph in (node for node in root.iter() if node.tag.endswith("}p")):
        value = "".join(node.text or "" for node in paragraph.iter() if node.tag.endswith("}t")).strip()
        if value:
            paragraphs.append(value)
    return "\n".join(paragraphs)


def load_docx(path: Path) -> str:
    try:
        from docx import Document

        return "\n".join(paragraph.text for paragraph in Document(str(path)).paragraphs)
    except (ImportError, KeyError, ValueError):
        return _load_docx_xml(Path(path))


def load_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    lines: list[str] = []
    try:
        for sheet in workbook.worksheets:
            lines.append(f"【工作表】{sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
                if values:
                    lines.append(" | ".join(values))
    finally:
        workbook.close()
    return "\n".join(lines)


def load_image(path: Path) -> str:
    from PIL import Image, ImageOps
    import pytesseract

    try:
        with Image.open(path) as image:
            image = ImageOps.exif_transpose(image).convert("L")
            if min(image.size) < 1200:
                image = image.resize((image.width * 2, image.height * 2))
            return pytesseract.image_to_string(image, lang="chi_sim+eng", config="--psm 6").strip()
    except Exception as exc:
        raise RuntimeError(f"图片 OCR 解析失败：{exc}") from exc


def load(path: Path) -> str:
    source = Path(path)
    suffix = source.suffix.lower()
    loaders = {".json": load_json, ".pdf": load_pdf, ".docx": load_docx, ".xlsx": load_xlsx}
    if suffix in TEXT_SUFFIXES:
        return loaders.get(suffix, load_text)(source)
    if suffix in IMAGE_SUFFIXES:
        return load_image(source)
    if suffix in loaders:
        return loaders[suffix](source)
    raise ValueError(f"不支持的文件类型：{suffix or '无后缀'}")


def extract_content(path: Path) -> ExtractedDocument:
    source = Path(path)
    text = load(source)
    suffix = source.suffix.lower()
    image = suffix in IMAGE_SUFFIXES
    pdf_ocr = suffix == ".pdf" and "【PDF OCR" in text
    method = "ocr" if image else ("ocr_pdf" if pdf_ocr else ("json" if suffix == ".json" else "text"))
    pages = [{"page": 1, "text": text}] if text else []
    return ExtractedDocument(text=text, pages=pages, extract_method=method, ocr_used=image or pdf_ocr, vision_used=False)


extract = load


__all__ = [
    "ExtractedDocument", "IMAGE_SUFFIXES", "SUPPORTED_SUFFIXES", "TEXT_SUFFIXES",
    "extract", "extract_content", "load", "load_docx", "load_image", "load_json",
    "load_pdf", "load_text", "load_xlsx",
]
