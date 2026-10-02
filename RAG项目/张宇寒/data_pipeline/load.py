"""根据文件类型选择读取技术，把原始资料安全转换为文字和分页信息。

普通文本直接读取；DOCX 用 python-docx；XLSX 用 openpyxl；文字 PDF 用 pypdf 和
pdfplumber；图片或扫描 PDF 先 OCR，用户明确允许时才用多模态模型补读。

``load(path)`` 只返回文字以兼容旧代码；``extract_content(path)`` 还会返回分页、
表格、读取方式和失败原因。这里不切块、不调用嵌入模型，也不连接数据库。
"""

# 延迟解析类型注解。
from __future__ import annotations

# closing 保证 PIL 图片等资源即使发生异常也会被关闭。
from contextlib import closing
# dataclass 简化读取结果类；asdict 转字典；field 为列表创建独立默认值。
from dataclasses import asdict, dataclass, field
# lru_cache 让本地 OCR 模型只初始化一次并复用。
from functools import lru_cache
# sha256 为文件内容生成稳定指纹。
from hashlib import sha256
# json 读取和重新序列化 JSON 文件。
import json
# os 读取 OCR 程序路径和语言环境变量。
import os
# Path 统一文件路径操作。
from pathlib import Path
# Any 表示 JSON 可以返回字典、列表、字符串、数字等任意合法结构。
from typing import Any


# 可以直接按文本读取的扩展名。
TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".rst", ".csv", ".tsv", ".log", ".json"}
# 交给 PIL 和 OCR 处理的常见图片扩展名。
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}
# 项目允许读取的全部扩展名。
SUPPORTED_SUFFIXES = TEXT_SUFFIXES | IMAGE_SUFFIXES | {".pdf", ".docx", ".xlsx"}
# 单文件上限 50MB，避免上传或解析超大文件耗尽内存。
MAX_FILE_BYTES = 50 * 1024 * 1024
# PDF 和多帧图片都限制为 500 页/帧。
MAX_PDF_PAGES = 500
MAX_IMAGE_FRAMES = 500


@dataclass
class ExtractedDocument:
    """统一描述任意格式文件的读取结果。"""

    # 合并后的全文。
    text: str
    # 每页、每帧或每工作表的文字与结构化信息。
    pages: list[dict]
    # 实际主要读取技术，例如 pypdf、rapidocr、vision。
    extract_method: str
    # 是否使用 OCR；None 表示外部工具没有提供该信息。
    ocr_used: bool | None
    # 是否把页面发送给多模态模型；None 同样表示未知。
    vision_used: bool | None
    # success / partial / failed / empty 描述整体读取完整性。
    status: str = "success"
    # warnings 保存失败页、公式未计算、混合图文需复核等提示。
    warnings: list[str] = field(default_factory=list)

    def to_dict(self) -> dict:
        """把 dataclass 转为普通字典，便于返回 API 或保存 JSON。"""

        return asdict(self)


def extract_content(path: Path, *, pdf_method: str = "auto", encoding: str = "utf-8-sig") -> ExtractedDocument:
    """按文件后缀分发读取器，并返回包含完整状态的 ExtractedDocument。"""

    # 把字符串路径和 Path 统一为 Path。
    source = Path(path)
    # 必须是实际文件；目录或不存在路径不能继续读取。
    if not source.is_file():
        raise FileNotFoundError(f"找不到文件：{source}")
    # 超大文件要求用户先拆分，避免解析过程占用过多内存和时间。
    if source.stat().st_size > MAX_FILE_BYTES:
        raise ValueError("文件超过50MB，请拆分后读取。")
    # 只接受三种明确策略，禁止拼写错误悄悄改变隐私行为。
    if pdf_method not in {"auto", "auto-vision", "vision"}:
        raise ValueError("pdf_method只能是auto（本地）、auto-vision（OCR后补读）或vision（直接多模态）。")
    # 扩展名统一成小写，兼容 .PDF、.JPG 等写法。
    suffix = source.suffix.lower()
    # 未支持格式不尝试猜测，直接给出明确错误。
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"不支持的文件类型：{suffix or '无后缀'}")
    # PDF 根据参数选择本地、OCR 补读或直接多模态路线。
    if suffix == ".pdf":
        return read_pdf(source, use_vision=pdf_method == "vision", allow_vision=pdf_method == "auto-vision")
    # Word 和 Excel 分别交给专用结构化读取器。
    if suffix == ".docx":
        return read_docx(source)
    if suffix == ".xlsx":
        return read_xlsx(source)
    # 图片与 PDF 扫描页复用 OCR/多模态规则。
    if suffix in IMAGE_SUFFIXES:
        return read_image(source, use_vision=pdf_method == "vision", allow_vision=pdf_method == "auto-vision")
    # JSON 先解析验证格式，再转回中文 JSON 文本供后续解析。
    if suffix == ".json":
        text = json.dumps(load_json(source), ensure_ascii=False)
        method = "json"
    else:
        # 其他文本格式按指定编码直接读取。
        text = load_text(source, encoding=encoding)
        method = "text"
    # 普通文本没有真实分页，用 page=None 表明页码未知而不是第一页。
    return ExtractedDocument(text, [{"page": None, "text": text}], method, False, False,
                             status="success" if text.strip() else "empty")


def load(path: Path, *, pdf_method: str = "auto", encoding: str = "utf-8-sig") -> str:
    """兼容旧流程：读取文件并只返回完整文字，不完整时直接报错。"""

    # extract_content 保留详细状态，complete_text 再保证结果可安全继续处理。
    return complete_text(extract_content(path, pdf_method=pdf_method, encoding=encoding))


def load_text(path: Path, *, encoding: str = "utf-8-sig") -> str:
    """默认UTF-8（包括BOM）；其他编码明确指定，不能悄悄丢字。"""

    # read_text 遇到错误编码会抛异常，避免使用 errors=ignore 丢失法律文字。
    return Path(path).read_text(encoding=encoding)


def load_json(path: Path) -> Any:
    """保留JSON的字典、列表结构；格式错误直接报错。"""

    # 先按统一文本方式读取，再由 json.loads 严格解析。
    return json.loads(load_text(path))


def file_sha256(path: str | Path) -> str:
    """分块读取文件并计算 SHA-256 内容指纹。"""

    # 创建哈希计算器。
    digest = sha256()
    # 二进制读取保证任意格式都按原始字节计算。
    with Path(path).open("rb") as file:
        # 每次读取 1MB；iter 遇到 b"" 表示文件结束。
        for block in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(block)
    # 十六进制字符串便于作为目录名或去重标识。
    return digest.hexdigest()


def table_text(rows: list[list]) -> str:
    """保留空格子的位置，不把后面的内容向前挤。"""

    # lines 保存每一行转成的纯文本。
    lines = []
    # 逐行处理表格。
    for row in rows:
        # None 表示空单元格；单元格内换行转为空格，避免破坏一行的列结构。
        cells = ["" if cell is None else str(cell).strip().replace("\n", " ") for cell in row]
        # 用竖线保留列之间的可见边界。
        lines.append(" | ".join(cells))
    return "\n".join(lines)


def read_docx(path: Path) -> ExtractedDocument:
    """按 Word 原始顺序读取段落和表格。"""

    # python-docx 负责打开 DOCX；Paragraph/Table 用于区分块类型。
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    # 打开 Word 文档。
    document = Document(str(path))
    # blocks 记录段落和表格，保留它们在文档中的先后顺序。
    blocks = []
    # iter_inner_content 会按原顺序给出正文、表格、后续正文。
    for item in document.iter_inner_content():
        # 非空段落直接保存其文字。
        if isinstance(item, Paragraph) and item.text.strip():
            blocks.append({"type": "paragraph", "text": item.text.strip()})
        # 表格逐行逐格读取，同时保留结构化 rows 和纯文本版本。
        elif isinstance(item, Table):
            rows = [[cell.text for cell in row.cells] for row in item.rows]
            blocks.append({"type": "table", "rows": rows, "text": table_text(rows)})
    # 用换行按原顺序连接所有块，供后续统一解析。
    text = "\n".join(block["text"] for block in blocks)
    # python-docx 没有渲染真实页面，因此 page 必须是 None，不能虚构页码。
    return ExtractedDocument(text, [{"page": None, "text": text, "blocks": blocks}], "python-docx", False, False,
                             status="success" if text.strip() else "empty")


def load_docx(path: Path) -> str:
    """旧接口：读取 Word 并只返回文字。"""

    return read_docx(path).text


def read_xlsx(path: Path) -> ExtractedDocument:
    """逐工作表读取 Excel 单元格地址、值和公式提示。"""

    # openpyxl 负责读取工作簿，get_column_letter 把数字列转为 A/B/C 地址。
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    # read_only 降低大表内存占用；data_only=False 保留公式原文，不伪装成计算结果。
    workbook = load_workbook(path, read_only=True, data_only=False)
    # sheets 保存每张工作表的文字和结构。
    sheets = []
    # warnings 记录需要用户核对的公式单元格。
    warnings = []
    # finally 保证中途异常也会关闭工作簿。
    try:
        # 按工作簿原顺序遍历工作表。
        for sheet in workbook.worksheets:
            # lines 是面向后续文本解析的行，rows 是保留地址的结构化行。
            lines = []
            rows = []
            # 从第 1 行开始遍历单元格。
            for row_number, row in enumerate(sheet.iter_rows(), start=1):
                # 取出这一行全部值。
                values = [cell.value for cell in row]
                # 完全空白的行没有意义，跳过。
                if all(value is None for value in values):
                    continue
                # cells 保存当前非空行的地址和值。
                cells = []
                # 从第 1 列开始逐格处理。
                for column, value in enumerate(values, start=1):
                    # 组合出 A1、B2 等人类可读地址。
                    address = f"{get_column_letter(column)}{row_number}"
                    # 空值保留为空字符串，不能让后面单元格向前错位。
                    cells.append({"address": address, "value": "" if value is None else str(value)})
                    # data_type == "f" 表示公式；读取器不会自行计算，必须提示核对。
                    if row[column - 1].data_type == "f":
                        warnings.append(f"工作表{sheet.title}的{address}是公式，未计算，请核对结果。")
                # 纯文本同时保留单元格地址，减少表格转文本后的歧义。
                lines.append(" | ".join(f"{cell['address']}={cell['value']}" for cell in cells))
                rows.append({"row": row_number, "cells": cells})
            # 合并当前工作表的文本行。
            text = "\n".join(lines)
            # page=None 表示这是工作表而不是真实页码。
            sheets.append({"page": None, "sheet": sheet.title, "text": text, "rows": rows})
    finally:
        # 释放文件句柄。
        workbook.close()
    # 不同工作表之间加入标题，避免后续解析混淆来源。
    text = "\n".join(f"【工作表】{sheet['sheet']}\n{sheet['text']}" for sheet in sheets)
    # 至少一张表存在有效行才算有内容。
    has_content = any(sheet["rows"] for sheet in sheets)
    # Excel 未使用 OCR 或多模态；公式提示放入 warnings。
    return ExtractedDocument(text, sheets, "openpyxl", False, False,
                             status="success" if has_content else "empty", warnings=warnings)


def load_xlsx(path: Path) -> str:
    """旧接口：读取 Excel 并只返回合并文字。"""

    return read_xlsx(path).text


@lru_cache(maxsize=1)
def rapidocr_engine():
    """复用项目已有的RapidOCR包；模型只初始化一次。"""

    # 延迟导入使不处理图片的任务无需加载 OCR 依赖。
    from rapidocr_onnxruntime import RapidOCR
    # 第一次调用创建模型，之后由 lru_cache 返回同一实例。
    return RapidOCR()


def read_image_with_rapidocr(image) -> str:
    """用本地 RapidOCR 识别一张 PIL 图片。"""

    # RapidOCR 需要 NumPy 图像数组。
    import numpy as np
    # PIL 转 RGB 后再反转颜色通道，得到 OpenCV 常用的 BGR 顺序。
    pixels = np.array(image.convert("RGB"))[:, :, ::-1].copy()
    # RapidOCR 返回识别结果和耗时等附加数据，此处只需要 result。
    result, _ = rapidocr_engine()(pixels)
    # 每个识别项的第 2 个值是文字；过滤空结果并按行连接。
    return "\n".join(str(item[1]).strip() for item in (result or []) if str(item[1]).strip())


def read_image_with_tesseract(image) -> str:
    """在 RapidOCR 无结果时，用本地 Tesseract 作为第二种 OCR。"""

    import pytesseract
    # Windows 上可通过环境变量明确指定 tesseract.exe 路径。
    if os.getenv("TESSERACT_CMD"):
        pytesseract.pytesseract.tesseract_cmd = os.environ["TESSERACT_CMD"]
    # 默认识别中英文，PSM 6 假设页面是一整块文字，单页最多等待 60 秒。
    return pytesseract.image_to_string(image, lang=os.getenv("OCR_LANGUAGE", "chi_sim+eng"),
                                      config="--psm 6", timeout=60).strip()


def ocr_image(image) -> tuple[str, str]:
    """先RapidOCR；失败或没有文字，再尝试Tesseract。"""

    # errors 收集每个 OCR 失败原因，最终一起返回便于排查环境。
    errors = []
    # 按轻量 RapidOCR、Tesseract 后备的顺序尝试。
    for name, read in [("rapidocr", read_image_with_rapidocr), ("tesseract", read_image_with_tesseract)]:
        try:
            # 调用当前 OCR 并清理首尾空白。
            text = read(image).strip()
            # 任一方法成功识别非空文字就立即返回文字和方法名。
            if text:
                return text, name
            errors.append(f"{name}没有识别出文字")
        except Exception as exc:
            # 单个 OCR 缺依赖或运行失败时继续尝试下一个。
            errors.append(f"{name}：{exc}")
    # 所有本地 OCR 都失败时，由上层决定是否允许多模态补读。
    raise RuntimeError("OCR读取失败：" + "；".join(errors))


def read_image_with_vision(image) -> str:
    """复用现有多模态配置；调用者必须已明确选择云端读取。"""

    # base64 把图片编码到 JSON 请求；BytesIO 在内存中暂存 PNG。
    import base64
    from io import BytesIO
    # httpx 发送多模态模型 HTTP 请求。
    import httpx
    # 从项目配置读取接口地址、密钥、模型名、超时和输出上限。
    from backend.app.config import get_settings

    # 加载当前项目配置。
    settings = get_settings()
    # 没有密钥时禁止发送，给出明确配置提示。
    if not settings.siliconflow_api_key:
        raise RuntimeError("多模态读取需要配置SILICONFLOW_API_KEY。")
    # 将 PIL 图片保存为内存 PNG，再进行 base64 编码。
    with BytesIO() as buffer:
        image.save(buffer, format="PNG")
        encoded = base64.b64encode(buffer.getvalue()).decode("ascii")
    # 提示词只允许忠实转写，不允许模型总结、建议或猜测看不清内容。
    prompt = ("逐行提取这页资料的原文，保留标题、条文编号、金额、日期和表格行列。"
              "不要总结，不要提供建议，不要补写或改写原文；看不清处写【看不清】。"
              "只返回提取内容，不要添加Markdown代码围栏。")
    # 调用与 OpenAI Chat Completions 格式兼容的多模态接口。
    response = httpx.post(
        # rstrip 避免配置末尾斜杠与路径拼接成双斜杠。
        settings.siliconflow_base_url.rstrip("/") + "/chat/completions",
        # Bearer Token 用于服务端鉴权。
        headers={"Authorization": f"Bearer {settings.siliconflow_api_key}"},
        # messages 中同时放入文字指令和 data URL 图片。
        json={"model": settings.multimodal_model, "messages": [{"role": "user", "content": [
            {"type": "text", "text": prompt},
            {"type": "image_url", "image_url": {"url": "data:image/png;base64," + encoded}}]}],
              # 温度为 0 降低改写随机性；max_tokens 控制最大输出长度。
              "temperature": 0, "max_tokens": settings.multimodal_max_tokens},
        # 请求超过配置时间仍未完成就终止。
        timeout=settings.multimodal_timeout,
    )
    # 4xx/5xx 直接转成异常，不能把错误页面当正文。
    response.raise_for_status()
    # 读取第一条模型候选结果。
    choice = response.json()["choices"][0]
    # 取出模型文字并清理首尾空白。
    text = str(choice["message"].get("content") or "").strip()
    # 非 stop 通常表示截断；空文、乱码或“看不清”都不能视为完整法律资料。
    if choice.get("finish_reason") != "stop" or not text or any(mark in text for mark in ("\ufffd", "【看不清】")):
        raise RuntimeError("多模态结果未完整识别，请检查原文件或调整模型输出上限。")
    return text


def read_scanned_image(image, *, use_vision: bool = False, allow_vision: bool = False, reference_text: str = "") -> tuple[str, str]:
    """图片和PDF共用：直接多模态，或先OCR、必要时补读。"""

    # use_vision 是用户明确选择直接云端多模态，不先运行本地 OCR。
    if use_vision:
        return read_image_with_vision(image), "vision"
    # 混合 PDF 可能已有参考文字层；识别结果低于其非空白字符数 80% 时怀疑漏读。
    # reference_text 为空时 minimum_length 为 0，不会用短文本字数误判普通图片。
    minimum_length = len("".join(reference_text.split())) * 0.8
    # 先尝试完全本地的 OCR，隐私和成本都更可控。
    try:
        text, method = ocr_image(image)
        text = text.strip()
        # 空结果、乱码和看不清标记都视为失败。
        if not text or any(mark in text for mark in ("\ufffd", "【看不清】")):
            raise RuntimeError("OCR结果为空、含乱码或看不清标记。")
        # 有参考文字层时，本地 OCR 明显更短可能漏掉原文。
        if len("".join(text.split())) < minimum_length:
            raise RuntimeError("OCR结果比原有文字层明显更少，可能漏读。")
        # 返回识别文字和实际成功的 OCR 名称。
        return text, method
    except Exception:
        # auto 模式禁止自动上传，直接保留 OCR 异常。
        if not allow_vision:
            raise
    # auto-vision 明确允许时，OCR 失败才调用多模态补读。
    text = read_image_with_vision(image)
    # 多模态结果同样与已有文字层做基本长度对照。
    if len("".join(text.split())) < minimum_length:
        raise RuntimeError("多模态结果比原有文字层明显更少，请人工核对。")
    return text, "vision"


def read_image(path: Path, *, use_vision: bool = False, allow_vision: bool = False) -> ExtractedDocument:
    """读取单张或多帧图片，并为每帧记录文字、方法和状态。"""

    # PIL 打开图片；ImageOps.exif_transpose 根据拍摄方向自动旋转。
    from PIL import Image, ImageOps
    # pages 中每一项对应图片的一帧。
    pages = []
    # warnings 收集失败帧原因。
    warnings = []
    # with 保证源图片文件最终关闭。
    with Image.open(path) as source:
        # 普通图片为 1，GIF/TIFF 等可能包含多帧。
        frames = getattr(source, "n_frames", 1)
        # 防止异常多帧文件占用过多时间和内存。
        if frames > MAX_IMAGE_FRAMES:
            raise ValueError("图片超过500帧，请拆分后读取。")
        # 按顺序读取每一帧。
        for index in range(frames):
            # 先设置默认值，即使读取失败也能生成结构完整的页面记录。
            text = ""
            method = "ocr"
            status = "success"
            try:
                # 将 PIL 当前帧移动到 index。
                source.seek(index)
                # exif_transpose 返回的新图像用 closing 保证释放像素内存。
                with closing(ImageOps.exif_transpose(source)) as image:
                    # 先给出计划方法，真正成功后会被 read_scanned_image 返回的方法覆盖。
                    method = "vision" if use_vision else ("ocr+vision" if allow_vision else "ocr")
                    text, method = read_scanned_image(image, use_vision=use_vision, allow_vision=allow_vision)
            except Exception as exc:
                # 单帧失败时继续记录其他帧，但整体结果会标为 partial/failed。
                status = "failed"
                warnings.append(f"图片第{index + 1}帧读取失败：{exc}")
            # 页码从 1 开始对用户更直观。
            pages.append({"page": index + 1, "text": text, "method": method, "status": status})
    # 单张图片直接使用正文；多帧文件增加帧标题，防止内容边界丢失。
    text = pages[0]["text"] if len(pages) == 1 else "\n\n".join(
        f"【图片第{page['page']}帧】\n{page['text']}" for page in pages if page["text"])
    # 任一帧失败都说明结果不完整。
    incomplete = any(page["status"] == "failed" for page in pages)
    # 有部分文字为 partial；全部失败为 failed；全部成功为 success。
    status = "partial" if incomplete and text else ("failed" if incomplete else "success")
    # 单帧保留真实方法，多帧用统一名称表示图片 OCR 流程。
    method = pages[0]["method"] if len(pages) == 1 else "ocr_image"
    # 只要任一帧用了 vision，就必须在结果中明确标记。
    vision_used = any("vision" in page["method"] for page in pages)
    # not use_vision 表示流程尝试过 OCR；auto-vision 后备也属于先 OCR。
    return ExtractedDocument(text, pages, method, not use_vision, vision_used, status=status, warnings=warnings)


def load_image(path: Path) -> str:
    """旧接口：本地读取图片并只在结果完整时返回文字。"""

    return complete_text(read_image(path))


def render_pdf_page(path: Path, page: int):
    """仅渲染需要OCR的这一页，不一次加载整本扫描PDF。"""

    # pdfplumber 用于把指定页渲染成图片。
    import pdfplumber
    # 每次只打开并渲染当前页，220 DPI 在识别清晰度和内存之间折中。
    with pdfplumber.open(path) as pdf:
        image = pdf.pages[page - 1].to_image(resolution=220).original
        # copy 后返回独立图像，PDF 关闭后仍可供 OCR 使用。
        return image.copy()


def read_pdf_tables(page) -> tuple[list[dict], str | None]:
    """文字PDF用pdfplumber找表格；没有表格时使用pypdf正文。"""

    # tables 保存表格坐标、单元格和纯文本。
    tables = []
    # find_tables 根据页面线条和文字布局发现表格区域。
    for table in page.find_tables():
        # extract 把表格转成二维行列数组。
        rows = table.extract()
        # bbox 是表格在页面上的左、上、右、下坐标。
        tables.append({"type": "table", "bbox": list(table.bbox), "rows": rows, "text": table_text(rows)})
    # 没有表格时返回 None，调用者继续使用 pypdf 已取得的正文。
    if not tables:
        return [], None

    def outside_tables(obj):
        """判断页面对象是否位于所有表格之外。"""

        # 非字符对象不属于需要排除的正文字符，保留。
        if obj.get("object_type") != "char":
            return True
        # 用字符中心点判断它是否落在某个表格边界内。
        x = (obj["x0"] + obj["x1"]) / 2
        y = (obj["top"] + obj["bottom"]) / 2
        # 不在任何表格 bbox 中才算表格外正文。
        return not any(t["bbox"][0] <= x <= t["bbox"][2] and t["bbox"][1] <= y <= t["bbox"][3] for t in tables)

    # 排除表格内字符后提取普通正文，避免同一金额既出现在正文又出现在表格中。
    lines = page.filter(outside_tables).extract_text_lines(return_chars=False)
    # 普通正文块保存文字和纵向坐标。
    blocks = [{"text": line["text"], "top": line["top"]} for line in lines]
    # 表格文本也作为块加入，并使用表格顶部坐标。
    blocks += [{"text": table["text"], "top": table["bbox"][1]} for table in tables]
    # 按页面从上到下重新排序，使表格回到原本出现位置。
    blocks.sort(key=lambda block: block["top"])
    # 同时返回结构化表格和按页面顺序组合的正文。
    return tables, "\n".join(block["text"] for block in blocks)


def read_pdf(path: Path, *, use_vision: bool = False, allow_vision: bool = False) -> ExtractedDocument:
    """逐页读取 PDF：文字页走 pypdf/pdfplumber，扫描页走 OCR 或多模态。"""

    # pypdf 读取文字层、图片和绘制操作。
    from pypdf import PdfReader
    # pdfplumber 识别表格并保留页面布局顺序。
    import pdfplumber

    # 打开 PDF 基础结构。
    pdf = PdfReader(str(path))
    # 空密码也无法解密时，要求用户先提供可读取版本。
    if pdf.is_encrypted and not pdf.decrypt(""):
        raise ValueError("PDF已加密，请先提供可读取的文件。")
    # 超过页数上限时要求拆分，避免一次任务过重。
    if len(pdf.pages) > MAX_PDF_PAGES:
        raise ValueError(f"PDF超过{MAX_PDF_PAGES}页，请拆分后读取。")
    # pages 保存每页文字、技术、表格和状态。
    pages = []
    # warnings 收集失败页和需要人工复核的混合图文页。
    warnings = []
    # ocr_used 记录是否至少一页使用本地 OCR。
    ocr_used = False
    # layout 与 pypdf 同时打开同一文件，用于表格处理。
    with pdfplumber.open(path) as layout:
        # 页码从 1 开始，便于用户定位原文。
        for number, page in enumerate(pdf.pages, start=1):
            # 每页先设置安全默认值。
            text = ""
            method = "pypdf"
            tables = []
            status = "success"
            try:
                # pypdf 优先提取已有文字层，成本最低且不会上传数据。
                text = (page.extract_text() or "").strip()
                # page.images 判断页面是否含位图。
                has_image = bool(page.images)
                # PDF 内容流可用于判断页面是否真正绘制了内容。
                content = page.get_contents()
                # 只把实际绘制路径、图片、阴影或文字的操作视为页面内容。
                # q/Q 等只调整图形状态，不能据此判断需要 OCR。
                paint_operations = {b"S", b"s", b"f", b"F", b"f*", b"B", b"B*", b"b", b"b*",
                                    b"Do", b"sh", b"Tj", b"TJ", b"'", b'"'}
                # 内容流存在且至少有一个真实绘制操作时为 True。
                has_drawing = content is not None and any(operation in paint_operations for _, operation in content.operations)
                # 乱码、没有文字但有绘制内容，或图片页文字层过短时需要 OCR。
                # 字数只是辅助条件，短的纯文字页面不会因短而自动 OCR。
                needs_ocr = "\ufffd" in text or (not text and has_drawing) or (has_image and len(text) < 80)
                # 明确 vision 会处理所有有绘制内容页；auto-vision 对含图片页允许补读。
                if needs_ocr or (use_vision and has_drawing) or (allow_vision and has_image):
                    # 直接 vision 不算 OCR；其余扫描路线先尝试过 OCR。
                    ocr_used = ocr_used or not use_vision
                    # 只渲染当前页，并确保识别后关闭图片。
                    with closing(render_pdf_page(path, number)) as image:
                        # 先记录计划方法，实际成功方法会由 read_scanned_image 覆盖。
                        method = "vision" if use_vision else ("ocr+vision" if allow_vision else "ocr")
                        # reference_text 用来发现 OCR 是否比原文字层明显漏字。
                        text, method = read_scanned_image(image, use_vision=use_vision, allow_vision=allow_vision, reference_text=text)
                else:
                    # 普通文字页检查表格，避免列内容顺序错乱。
                    tables, table_body = read_pdf_tables(layout.pages[number - 1])
                    # 找到表格时改用融合了正文和表格顺序的文本。
                    if table_body is not None:
                        text = table_body
                        method = "pypdf+pdfplumber"
                    # 有图片但未走视觉识别，可能遗漏图片文字，因此标记需人工复核。
                    if has_image:
                        status = "needs_review"
                        warnings.append(f"第{number}页同时包含文字和图片；请人工复核，或明确选择vision读取。")
            except Exception as exc:
                # 单页失败不立即丢弃其他页，但整体状态会是 partial/failed。
                status = "failed"
                warnings.append(f"第{number}页未完整读取：{exc}")
            # 无论成功失败都保存页面记录，方便精确定位问题。
            pages.append({"page": number, "text": text, "method": method, "tables": tables, "status": status})

    # 合并所有非空页，并加页码标签供预览；解析公共法律正文时会使用 pages 原文避免标签混入。
    text = "\n\n".join(f"【PDF第{page['page']}页】\n{page['text']}" for page in pages if page["text"])
    # failed 和 needs_review 都意味着不能直接认定完全读取。
    incomplete = any(page["status"] != "success" for page in pages)
    # 根据是否不完整、是否还有部分文字，确定整体状态。
    status = "partial" if incomplete and text else ("failed" if incomplete else ("success" if text else "empty"))
    # 任一页方法名包含 vision，就明确标记发生了云端多模态读取。
    vision_used = any("vision" in page["method"] for page in pages)
    # extract_method 概括整份 PDF 的最高级读取路线。
    return ExtractedDocument(text, pages, "vision" if vision_used else ("ocr_pdf" if ocr_used else "pypdf+pdfplumber"), ocr_used, vision_used,
                             status=status, warnings=warnings)


def complete_text(result: ExtractedDocument) -> str:
    """旧接口只返回文字，因此不能把失败信息藏起来。"""

    # partial/failed 说明至少有页面未正确读取，抛出全部警告而不是返回残缺正文。
    if result.status in {"partial", "failed"}:
        raise RuntimeError("文件未完整读取：" + "；".join(result.warnings))
    # 空文件也不能进入解析和向量库。
    if result.status == "empty":
        raise RuntimeError("文件没有可读取的内容，请检查原文件。")
    # 只有完整成功时才返回正文。
    return result.text


def load_pdf(path: Path) -> str:
    """旧接口：使用默认本地策略读取 PDF，并只返回完整文字。"""

    return complete_text(read_pdf(path))


# extract 是旧代码使用的名称，继续指向统一 load 入口。
extract = load

# 明确本模块允许外部导入的类型、常量和入口函数。
__all__ = [
    "ExtractedDocument", "IMAGE_SUFFIXES", "SUPPORTED_SUFFIXES", "TEXT_SUFFIXES",
    "extract", "extract_content", "file_sha256", "load", "load_docx", "load_image",
    "load_json", "load_pdf", "load_text", "load_xlsx",
]
