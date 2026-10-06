# md→docx 往返是 MinerU 不支持 md 输入的唯一出路，必须保证行内容零改写
import docx
from app.ingest.to_docx import md_to_docx


def test_roundtrip_preserves_every_line(tmp_path):
    src = tmp_path / "src.md"
    lines = ["# 标题", "", "## 第三编 合同", "", "第四百六十三条 正文内容。"]
    # 加一行含全角标点的正文，验证编码没被环境默认值吃掉
    lines.append("第一千零四十条 含特殊字符（）【】的正文。")
    src.write_text("\n".join(lines), encoding="utf-8")

    out = md_to_docx(src, tmp_path / "out.docx")
    got = [p.text for p in docx.Document(out).paragraphs]

    assert got == lines, "往返必须逐行原样，不得改写正文"


def test_empty_lines_become_empty_paragraphs(tmp_path):
    # 空行是款的分隔信号，丢了会让整条的款合并成一个块
    src = tmp_path / "src.md"
    src.write_text("第一千零四十一条 婚姻家庭受国家保护。\n\n实行婚姻自由。\n", encoding="utf-8")
    out = md_to_docx(src, tmp_path / "out.docx")
    got = [p.text for p in docx.Document(out).paragraphs]
    assert got == ["第一千零四十一条 婚姻家庭受国家保护。", "", "实行婚姻自由。"]


def test_creates_parent_directory(tmp_path):
    # 中册的合成件落在 data/interim/_synthetic/ 下，该目录初始不存在
    src = tmp_path / "src.md"
    src.write_text("第一条 内容。", encoding="utf-8")
    out = md_to_docx(src, tmp_path / "deep" / "nested" / "out.docx")
    assert out.exists()
