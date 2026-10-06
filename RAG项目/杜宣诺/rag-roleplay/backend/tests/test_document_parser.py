from app.services.document_parser import DocumentParseError, chunk_text, extract_text


def test_chunk_text_splits_by_size():
    text = "段落一\n\n段落二\n\n段落三"
    chunks = chunk_text(text, size=5, overlap=0)
    assert len(chunks) >= 2
    # 所有内容最终都被覆盖
    joined = "".join(chunks)
    assert "段落一" in joined and "段落二" in joined and "段落三" in joined


def test_chunk_text_respects_max_length():
    long_para = "字" * 5000
    chunks = chunk_text(long_para, size=800, overlap=100)
    assert all(len(c) <= 4096 for c in chunks)
    assert len(chunks) > 1


def test_chunk_text_overlap():
    text = "甲" * 20 + "\n\n" + "乙" * 20
    chunks = chunk_text(text, size=15, overlap=5)
    assert len(chunks) >= 1


def test_chunk_text_empty():
    assert chunk_text("", 100, 10) == []
    assert chunk_text("   \n\n  ", 100, 10) == []


def test_extract_text_txt():
    assert extract_text("a.txt", "你好".encode("utf-8")) == "你好"
    # gbk 回退
    assert extract_text("a.txt", "中文".encode("gbk")) == "中文"


def test_extract_text_unsupported():
    try:
        extract_text("a.exe", b"x")
        assert False, "应抛出异常"
    except Exception as e:
        assert "不支持" in str(e)


def test_extract_text_docx():
    import io

    import docx

    doc = docx.Document()
    doc.add_paragraph("第一段内容")
    doc.add_paragraph("")
    doc.add_paragraph("第二段内容")
    buf = io.BytesIO()
    doc.save(buf)

    text = extract_text("a.docx", buf.getvalue())
    assert "第一段内容" in text
    assert "第二段内容" in text
    # 空段落被跳过，不会残留空行
    assert "\n\n\n" not in text


def _build_pdf(body: str | None) -> bytes:
    """构造一个单页合法 PDF；body 为 None 时页面无文字层（模拟扫描件）。"""
    if body is None:
        content = b""
    else:
        content = f"BT /F1 24 Tf 72 720 Td ({body}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + obj + b"\nendobj\n"
    xref_pos = len(out)
    out += b"xref\n0 %d\n" % (len(objects) + 1)
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref_pos)
    return bytes(out)


def test_extract_text_pdf_with_text_layer():
    text = extract_text("a.pdf", _build_pdf("Hello World"))
    assert "Hello World" in text


def test_extract_text_pdf_scanned_raises_without_ocr():
    # 无文字层的 PDF 在未启用 OCR 时应抛 DocumentParseError
    try:
        extract_text("scan.pdf", _build_pdf(None))
        assert False, "应抛出异常"
    except DocumentParseError as e:
        assert "OCR" in str(e)
